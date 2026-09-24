"""XLSX report: Summary sheet + Findings sheet (conditional formatting by level, frozen header,
auto-filter, auto width) + Objects sheet. All labels in the requested language."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..core.models import Finding, ScanResult
from ..core.risk_engine import security_band
from . import COLUMN_KEYS, column_labels, finding_row
from .labels import band_name, category_name, labels, level_name, type_name

LEVEL_FILLS = {"Critical": "FECACA", "High": "FED7AA", "Medium": "FEF08A", "Low": "BBF7D0"}
HEADER_FILL = PatternFill("solid", fgColor="0F172A")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Border(bottom=Side(style="thin", color="CBD5E1"))
KPIS = ["inactive_users", "service_pne", "disabled_privileged", "excessive_rights", "hidden_admins", "kerberoastable",
        "asrep_roastable", "inactive_computers", "privileged", "service", "users", "computers"]


def _set(cell: Cell, value: object) -> None:
    """Write a value; directory-controlled strings are forced to text so '=...' never becomes a formula."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def _autowidth(ws, max_width: int = 60) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[letter].width = max(8, min(max_width, longest + 2))


def _level_rules(ws, rng: str, lang: str) -> None:
    for lvl, color in LEVEL_FILLS.items():
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$A2="{level_name(lvl, lang)}"'],
                                                       fill=PatternFill("solid", fgColor=color)))


def export_xlsx(scan: ScanResult, findings: list[Finding] | None = None, lang: str = "en") -> bytes:
    rows = findings if findings is not None else scan.findings
    lab = labels(lang)
    wb = Workbook()

    # ---- Summary -------------------------------------------------------------------------------
    ws = wb.active
    ws.title = lab["sheet.summary"]
    ws["A1"] = lab["report.title"]
    ws["A1"].font = Font(bold=True, size=16)
    meta = [
        (lab["meta.domain"], scan.domain), (lab["meta.scan_id"], scan.scan_id),
        (lab["meta.started"], scan.started_at.strftime("%Y-%m-%d %H:%M:%S")), (lab["meta.source"], scan.source),
        (lab["meta.trigger"], scan.trigger), (lab["meta.score"], scan.ad_security_score),
        (lab["meta.band"], band_name(security_band(scan.ad_security_score), lang)), (lab["meta.rows"], len(rows)),
    ]
    r = 3
    for k, v in meta:
        _set(ws.cell(row=r, column=1), k)
        ws.cell(row=r, column=1).font = Font(bold=True)
        _set(ws.cell(row=r, column=2), v)
        r += 1
    ws.cell(row=8, column=2).font = Font(bold=True, size=14)

    r += 1
    ws.cell(row=r, column=1, value=lab["sec.levels"]).font = Font(bold=True, size=12)
    r += 1
    for lvl in ("Critical", "High", "Medium", "Low"):
        ws.cell(row=r, column=1, value=level_name(lvl, lang)).fill = PatternFill("solid", fgColor=LEVEL_FILLS[lvl])
        ws.cell(row=r, column=2, value=scan.counts.get(lvl, 0))
        r += 1

    r += 1
    ws.cell(row=r, column=1, value=lab["sec.kpis"]).font = Font(bold=True, size=12)
    r += 1
    for key in KPIS:
        ws.cell(row=r, column=1, value=lab[f"kpi.{key}"])
        ws.cell(row=r, column=2, value=scan.counts.get(key, 0))
        r += 1

    r += 1
    ws.cell(row=r, column=1, value=lab["sec.categories"]).font = Font(bold=True, size=12)
    r += 1
    for cat, pen in scan.category_scores.items():
        ws.cell(row=r, column=1, value=category_name(cat, lang))
        ws.cell(row=r, column=2, value=pen)
        r += 1
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 44

    # ---- Findings ------------------------------------------------------------------------------
    wf = wb.create_sheet(lab["sheet.findings"])
    for c, label in enumerate(column_labels(lang), start=1):
        cell = wf.cell(row=1, column=c, value=label)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for i, f in enumerate(rows, start=2):
        row = finding_row(f, lang)
        for c, key in enumerate(COLUMN_KEYS, start=1):
            cell = wf.cell(row=i, column=c)
            _set(cell, row[key])
            cell.border = THIN
            if key in ("description", "recommendation", "remediation_command", "evidence"):
                cell.alignment = Alignment(wrap_text=False, vertical="top")
    wf.freeze_panes = "A2"
    last = max(2, len(rows) + 1)
    wf.auto_filter.ref = f"A1:{get_column_letter(len(COLUMN_KEYS))}{last}"
    _level_rules(wf, f"A2:{get_column_letter(len(COLUMN_KEYS))}{last}", lang)
    _autowidth(wf)

    # ---- Objects -------------------------------------------------------------------------------
    wo = wb.create_sheet(lab["sheet.objects"])
    head = [lab[k] for k in ("obj.level", "obj.score", "obj.k", "obj.type", "obj.object", "obj.dn", "obj.findings",
                             "obj.rules", "obj.path")]
    for c, label in enumerate(head, start=1):
        cell = wo.cell(row=1, column=c, value=label)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    ids = {f.object_id for f in rows}
    ents = [e for e in scan.entities if e.object_id in ids]
    for i, e in enumerate(ents, start=2):
        vals = [level_name(e.level.value, lang), e.score, e.k, type_name(e.object_type, lang), e.object_name, e.object_dn,
                e.finding_count, ", ".join(e.rule_ids), " → ".join(e.privilege_path or [])]
        for c, v in enumerate(vals, start=1):
            _set(wo.cell(row=i, column=c), v)
    wo.freeze_panes = "A2"
    _level_rules(wo, f"A2:A{max(2, len(ents) + 1)}", lang)
    _autowidth(wo)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
