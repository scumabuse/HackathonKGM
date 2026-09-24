"""CSV/XLSX/HTML files generate, open and are safe."""
import copy
import csv
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.exporters.csv_exporter import export_csv
from app.exporters.html_exporter import export_html
from app.exporters.xlsx_exporter import export_xlsx


def test_csv_opens_and_is_complete(mock_result):
    data = export_csv(mock_result)
    assert data.startswith("﻿".encode("utf-8"))  # Excel-friendly BOM
    rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
    assert len(rows) == len(mock_result.findings) > 0
    sa01 = [r for r in rows if r["Object"] == "ServiceAccount01"]
    assert len(sa01) == 5 and all(r["Object risk score"] == "85" for r in sa01)
    ivanov = next(r for r in rows if r["Rule"] == "AD-PRIV-NESTED" and r["Object"] == "ivanov")
    assert ivanov["Privilege path"] == "ivanov → IT-Support → Helpdesk-L2 → Domain Admins"


def test_xlsx_opens_with_summary_findings_and_formatting(mock_result):
    wb = load_workbook(io.BytesIO(export_xlsx(mock_result)))
    assert wb.sheetnames == ["Summary", "Findings", "Objects"]
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r[0]}
    assert summary["AD Security Score (0–100, higher is better)"] == mock_result.ad_security_score
    assert summary["Inactive users"] == 10 and summary["Excessive rights (nested groups)"] == 2
    ws = wb["Findings"]
    assert ws.freeze_panes == "A2"
    assert ws.max_row == len(mock_result.findings) + 1
    assert ws["A1"].value == "Level"
    assert len(ws.conditional_formatting) >= 1
    assert ws.column_dimensions["G"].width > 10  # auto width


def test_html_is_standalone_and_escaped(mock_result):
    html = export_html(mock_result).decode("utf-8")
    assert "<table>" in html and "corp.local" in html
    assert "ServiceAccount01" in html and "higher is better" in html
    for external in ("http://", "https://", "<script", "<link"):
        assert external not in html


def test_injection_is_neutralized(mock_result):
    evil = copy.deepcopy(mock_result)
    f = evil.findings[0]
    f.object_name = '=HYPERLINK("http://evil","x")'
    f.title = "<img src=x onerror=alert(1)>"
    csv_text = export_csv(evil, [f]).decode("utf-8-sig")
    assert "'=HYPERLINK" in csv_text
    html = export_html(evil, [f]).decode("utf-8")
    assert "<img src=x" not in html and "&lt;img src=x" in html
    wb = load_workbook(io.BytesIO(export_xlsx(evil, [f])))
    cell = wb["Findings"]["I2"]
    assert cell.value == '=HYPERLINK("http://evil","x")' and cell.data_type == "s"  # text, never a formula


@pytest.fixture()
def client(isolated_env):
    from app.main import create_app

    with TestClient(create_app()) as c:
        c.post("/api/scan", json={"source": "mock"})
        yield c


@pytest.mark.parametrize("fmt,ctype,magic", [("csv", "text/csv", b"\xef\xbb\xbf"), ("xlsx", "spreadsheetml", b"PK"),
                                              ("html", "text/html", b"<!doctype html>")])
def test_export_endpoint_downloads(client, fmt, ctype, magic):
    r = client.get("/api/export/latest", params={"format": fmt})
    assert r.status_code == 200
    assert ctype in r.headers["content-type"]
    assert f".{fmt}" in r.headers["content-disposition"]
    assert r.content.startswith(magic) and len(r.content) > 1000


def test_filtered_export_matches_filter(client):
    r = client.get("/api/export/latest", params={"format": "csv", "level": "Critical"})
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows and all(row["Level"] == "Critical" for row in rows)
    assert client.get("/api/export/latest", params={"format": "pdf"}).status_code == 422
