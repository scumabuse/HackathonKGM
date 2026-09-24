"""Localization: every rule/message exists in en/ru/kk, renders cleanly, and the API/exports honour ?lang=."""
import csv
import io
import re
import string

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.i18n import CATALOG, ago, tr
from app.core.rules_loader import RuleValidationError, load_rules, load_rules_from
from app.exporters.labels import LABELS

UNRESOLVED = re.compile(r"<[a-z_]+>|\{[a-z_]+\}")


def placeholders(text: str) -> set[str]:
    return {n for _, n, _, _ in string.Formatter().parse(text) if n}


def test_every_rule_is_fully_translated():
    for rid, rule in load_rules().items():
        for lang in ("ru", "kk"):
            tx = rule.i18n.get(lang)
            assert tx is not None, f"{rid} has no {lang} translation"
            for field in ("name", "title", "description", "recommendation"):
                assert getattr(tx, field), f"{rid}.{lang}.{field} is empty"


def test_translation_with_unknown_placeholder_is_rejected(tmp_path):
    from app.core.rules_loader import RULES_DIR

    for f in RULES_DIR.glob("*.yaml"):
        (tmp_path / f.name).write_text(f.read_text(encoding="utf-8"), encoding="utf-8")
    (tmp_path / "i18n").mkdir()
    (tmp_path / "i18n" / "ru.yaml").write_text('AD-USR-PNE:\n  title: "Пароль {typo_var}"\n', encoding="utf-8")
    with pytest.raises(RuleValidationError, match="typo_var"):
        load_rules_from(tmp_path)


def test_catalog_messages_have_consistent_placeholders():
    for key, (en, ru, kk) in CATALOG.items():
        assert en and ru and kk, key
        assert placeholders(en) == placeholders(ru) == placeholders(kk), key


def test_export_labels_have_identical_keys():
    assert LABELS["en"].keys() == LABELS["ru"].keys() == LABELS["kk"].keys()


def test_russian_plural_and_kazakh_ago():
    from datetime import UTC, datetime, timedelta

    now = datetime(2026, 9, 24, tzinfo=UTC)
    for n, word in [(1, "день"), (2, "дня"), (5, "дней"), (11, "дней"), (21, "день"), (730, "дней")]:
        assert ago(now - timedelta(days=n), now)["ru"].startswith(f"{n} {word} назад")
    assert ago(now - timedelta(days=730), now)["kk"].startswith("730 күн бұрын")
    assert ago(None, now) == tr("never")


def test_findings_render_cleanly_in_every_language(mock_result):
    for f in mock_result.findings:
        for lang in ("ru", "kk"):
            tx = f.i18n[lang]
            for field in ("title", "description", "recommendation"):
                assert tx[field] and not UNRESOLVED.search(tx[field]), (f.rule_id, lang, field, tx[field])
            assert tx["title"] != f.title, (f.rule_id, lang)


def test_evidence_and_facts_are_localized(mock_result):
    f = next(x for x in mock_result.findings if x.object_name == "q.mamyrov" and x.rule_id == "AD-USR-INACTIVE")
    assert f.evidence[0].value == "never logged on"
    assert f.evidence[0].value_i18n["ru"] == "ни разу не входил(а)"
    assert f.evidence[0].value_i18n["kk"] == "ешқашан кірмеген"
    e = next(x for x in mock_result.entities if x.object_name == "ServiceAccount01")
    assert e.attributes_i18n["ru"]["Статус"] == "Включена"
    assert e.attributes_i18n["kk"]["Құпиясөз ауыстырылған"].startswith("730 күн")
    assert e.top_title_i18n["ru"].startswith("Повышенные права")
    assert [w.title_i18n["ru"] for w in e.weight_breakdown if w.matched][0] == "Повышенные права"


@pytest.fixture()
def client(isolated_env):
    from app.main import create_app

    with TestClient(create_app()) as c:
        c.post("/api/scan", json={"source": "mock"})
        yield c


def test_api_answers_in_requested_language(client):
    ru = client.get("/api/findings", params={"lang": "ru", "rule_id": "AD-PRIV-NESTED"}).json()["items"][0]
    assert ru["title"].startswith("Избыточные права через вложенные группы")
    assert "i18n" not in ru and all("value_i18n" not in e for e in ru["evidence"])
    en = client.get("/api/findings", params={"rule_id": "AD-PRIV-NESTED"}).json()["items"][0]
    assert en["title"].startswith("Excessive rights via nested groups")
    kk = client.get(f"/api/accounts/{ru['object_id']}", params={"lang": "kk"}).json()
    assert kk["entity"]["top_title"].startswith("Кірістірілген топтар арқылы")
    assert "Күйі" in kk["entity"]["attributes"]
    dash = client.get("/api/dashboard/latest", params={"lang": "ru"}).json()
    sa01 = next(e for e in dash["top_risky"] if e["object_name"] == "ServiceAccount01")
    assert sa01["top_title"].startswith("Повышенные права")
    assert sa01["weight_breakdown"][0]["title"] == "Повышенные права"
    rules = client.get("/api/rules", params={"lang": "kk"}).json()
    assert next(r for r in rules if r["id"] == "AD-SVC-NO-OWNER")["name"] == "Иесі жоқ"
    settings = client.get("/api/settings", params={"lang": "ru"}).json()
    assert next(r for r in settings["rules"] if r["id"] == "AD-SVC-NO-OWNER")["name"] == "Нет владельца"
    assert client.get("/api/findings", params={"lang": "de"}).status_code == 422


def test_search_works_in_russian_and_kazakh(client):
    ru = client.get("/api/findings", params={"q": "НЕАКТИВНА", "lang": "ru"}).json()
    assert ru["total"] >= 10 and all(f["title"].lower().find("неактивна") >= 0 for f in ru["items"])
    kk = client.get("/api/findings", params={"q": "жасырын", "lang": "kk"}).json()
    assert [f["object_name"] for f in kk["items"]] == ["m.petrov"]


def test_exports_in_russian_and_kazakh(client):
    r = client.get("/api/export/latest", params={"format": "csv", "lang": "ru"})
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert "Уровень" in rows[0] and rows[0]["Уровень"] == "Критический"
    assert any(row["Находка"].startswith("Повышенные права") for row in rows)
    wb = load_workbook(io.BytesIO(client.get("/api/export/latest", params={"format": "xlsx", "lang": "kk"}).content))
    assert wb.sheetnames == ["Жиынтық", "Мәселелер", "Объектілер"]
    assert wb["Мәселелер"]["A2"].value == "Сыни"
    html = client.get("/api/export/latest", params={"format": "html", "lang": "ru"}).content.decode("utf-8")
    assert '<html lang="ru">' in html and "Отчёт о безопасности Active Directory" in html
    assert "чем выше, тем лучше" in html and "Критический" in html


def test_old_database_is_migrated(isolated_env):
    from sqlalchemy import create_engine, inspect, text

    from app.config import get_config
    from app.storage import db
    from app.storage.schema import Base

    url = get_config().DATABASE_URL
    eng = create_engine(url)
    Base.metadata.create_all(eng)
    with eng.begin() as conn:
        conn.execute(text("ALTER TABLE findings DROP COLUMN search_text"))
    assert "search_text" not in {c["name"] for c in inspect(eng).get_columns("findings")}
    eng.dispose()
    db.reset_engine()
    assert "search_text" in {c["name"] for c in inspect(db.get_engine()).get_columns("findings")}
